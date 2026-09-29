import os
import re
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import customtkinter as ctk

from ui.icons import icon
from ui.palette import COLORS, FONT, MONO, SEVERITY_COLORS, SEVERITY_ICONS
from ui.trend import TrendChart


class DetailModal(ctk.CTkToplevel):

    WIDTH = 1180
    HEIGHT = 720

    def __init__(self, parent, card, repository, perimetre="FXL", mode_execution="COMMIT"):
        super().__init__(parent)
        self.withdraw()  # on affiche la fenêtre une fois construite

        self.card = card
        self.repository = repository
        self.perimetre = perimetre
        self.mode_execution = mode_execution

        self.rows = []
        self.columns = []
        self.visible_values = []   # lignes actuellement affichées (pour l'export)

        self.title(f"CTRL-{card.control_id:02d} — {card.control_name}")
        self.configure(fg_color=COLORS["bg"])
        self.minsize(980, 600)
        self._place(parent)

        self.search_var = tk.StringVar()

        self._fonts()
        self._style_tree()
        self.build()
        self.load_data()

        self.transient(parent)
        self.deiconify()
        self.after(150, self._activate)

    # =============================================================
    # OUTILS FENETRE
    # =============================================================

    def _place(self, parent):
        try:
            parent.update_idletasks()
            x = parent.winfo_rootx() + (parent.winfo_width() - self.WIDTH) // 2
            y = parent.winfo_rooty() + (parent.winfo_height() - self.HEIGHT) // 2
            self.geometry(f"{self.WIDTH}x{self.HEIGHT}+{max(x, 0)}+{max(y, 0)}")
        except Exception:
            self.geometry(f"{self.WIDTH}x{self.HEIGHT}")

    def _activate(self):
        try:
            self.lift()
            self.focus_force()
            self.grab_set()
        except Exception:
            pass

    # =============================================================
    # POLICES / STYLE
    # =============================================================

    def _fonts(self):
        self.f_title = ctk.CTkFont(family=FONT, size=20, weight="bold")
        self.f_desc = ctk.CTkFont(family=FONT, size=11)
        self.f_id = ctk.CTkFont(family=MONO, size=10, weight="bold")
        self.f_btn = ctk.CTkFont(family=FONT, size=11, weight="bold")
        self.f_sum_label = ctk.CTkFont(family=FONT, size=10)
        self.f_sum_value = ctk.CTkFont(family=FONT, size=18, weight="bold")
        self.f_card_title = ctk.CTkFont(family=FONT, size=13, weight="bold")
        self.f_small = ctk.CTkFont(family=FONT, size=10)
        self.f_search = ctk.CTkFont(family=FONT, size=11)

    def _style_tree(self):

        style = ttk.Style(self)
        style.theme_use("clam")

        style.configure(
            "Detail.Treeview",
            background=COLORS["panel"], fieldbackground=COLORS["panel"],
            foreground=COLORS["text"], rowheight=36, borderwidth=0,
            font=(FONT, 10)
        )
        style.configure(
            "Detail.Treeview.Heading",
            background=COLORS["panel_2"], foreground=COLORS["muted"],
            font=(FONT, 9, "bold"), relief="flat", borderwidth=0,
            padding=(12, 9)
        )
        style.map(
            "Detail.Treeview",
            background=[("selected", COLORS["accent_soft"])],
            foreground=[("selected", COLORS["accent"])]
        )
        style.map(
            "Detail.Treeview.Heading",
            background=[("active", COLORS["track"])]
        )
        style.layout("Detail.Treeview", [("Treeview.treearea", {"sticky": "nswe"})])

    # =============================================================
    # BUILD
    # =============================================================

    def build(self):

        severite = self.card.severite or "N/A"
        color, soft = SEVERITY_COLORS.get(severite, SEVERITY_COLORS["N/A"])
        sev_icon = SEVERITY_ICONS.get(severite, "warning")

        # =========================================================
        # HEADER
        # =========================================================

        header = ctk.CTkFrame(self, fg_color=COLORS["panel"], corner_radius=0)
        header.pack(fill="x")

        ctk.CTkFrame(self, fg_color=COLORS["border"], height=1, corner_radius=0).pack(fill="x")

        inner = ctk.CTkFrame(header, fg_color="transparent")
        inner.pack(fill="x", padx=28, pady=20)

        box = ctk.CTkFrame(inner, fg_color=soft, width=52, height=52, corner_radius=14)
        box.pack(side="left", padx=(0, 16))
        box.pack_propagate(False)

        ctk.CTkLabel(
            box, text="", image=icon(sev_icon, color, 26), fg_color="transparent"
        ).pack(expand=True)

        titles = ctk.CTkFrame(inner, fg_color="transparent")
        titles.pack(side="left")

        line = ctk.CTkFrame(titles, fg_color="transparent")
        line.pack(anchor="w")

        ctk.CTkLabel(
            line, text=f"  CTRL-{self.card.control_id:02d}  ", font=self.f_id,
            text_color=COLORS["accent"], fg_color=COLORS["accent_soft"],
            corner_radius=8, height=24
        ).pack(side="left", padx=(0, 10))

        ctk.CTkLabel(
            line, text=self.card.control_name, font=self.f_title,
            text_color=COLORS["text"], fg_color="transparent"
        ).pack(side="left")

        ctk.CTkLabel(
            titles, text=self.card.description or "", font=self.f_desc,
            text_color=COLORS["muted"], fg_color="transparent",
            anchor="w", justify="left", wraplength=760
        ).pack(anchor="w", pady=(6, 0))

        ctk.CTkButton(
            inner, text="  Fermer", font=self.f_btn, command=self.destroy,
            image=icon("close", COLORS["text_soft"], 16), compound="left",
            fg_color=COLORS["panel"], hover_color=COLORS["panel_2"],
            text_color=COLORS["text_soft"], border_width=1,
            border_color=COLORS["border"], corner_radius=10, height=38, width=110
        ).pack(side="right")

        # =========================================================
        # SUMMARY
        # =========================================================

        summary = ctk.CTkFrame(self, fg_color="transparent")
        summary.pack(fill="x", padx=28, pady=(20, 14))

        for i in range(4):
            summary.grid_columnconfigure(i, weight=1, uniform="sum")

        montant = (
            f"{self.card.montant_impacte:,.0f}".replace(",", " ")
            if self.card.montant_impacte is not None
            else "—"
        )

        items = [
            (str(self.card.nb_items), "Éléments détectés", "list",
             COLORS["accent"], COLORS["accent_soft"], COLORS["text"]),
            (severite, "Sévérité", sev_icon, color, soft, color),
            (montant, "Montant impacté", "money",
             COLORS["accent"], COLORS["accent_soft"], COLORS["text"]),
            (self.card.frequence or "—", "Fréquence", "clock",
             COLORS["accent"], COLORS["accent_soft"], COLORS["text"]),
        ]

        for i, (value, label, icon_name, ico_color, ico_bg, val_color) in enumerate(items):
            self.create_summary(summary, i, value, label, icon_name,
                                ico_color, ico_bg, val_color)

        # =========================================================
        # BODY
        # =========================================================

        body = ctk.CTkFrame(self, fg_color="transparent")
        body.pack(fill="both", expand=True, padx=28, pady=(0, 24))

        # ---------------------------------------------------------
        # TABLE
        # ---------------------------------------------------------

        table_card = ctk.CTkFrame(
            body, fg_color=COLORS["panel"], corner_radius=16,
            border_width=1, border_color=COLORS["border"]
        )
        table_card.pack(side="left", fill="both", expand=True, padx=(0, 8))

        thead = ctk.CTkFrame(table_card, fg_color="transparent")
        thead.pack(fill="x", padx=20, pady=(16, 12))

        titles2 = ctk.CTkFrame(thead, fg_color="transparent")
        titles2.pack(side="left")

        ctk.CTkLabel(
            titles2, text="Éléments détectés", font=self.f_card_title,
            text_color=COLORS["text"], fg_color="transparent"
        ).pack(anchor="w")

        self.count_label = ctk.CTkLabel(
            titles2, text="Dernière exécution", font=self.f_small,
            text_color=COLORS["muted"], fg_color="transparent"
        )
        self.count_label.pack(anchor="w")

        # bouton d'export Excel (à droite, avant la recherche)
        self.export_btn = ctk.CTkButton(
            thead, text="Exporter Excel", font=self.f_btn,
            command=self.export_excel,
            fg_color=COLORS["accent"], hover_color=COLORS["accent"],
            text_color="#FFFFFF", corner_radius=10, height=38, width=130
        )
        self.export_btn.pack(side="right", padx=(10, 0))

        search = ctk.CTkFrame(
            thead, fg_color=COLORS["panel"], corner_radius=10,
            border_width=1, border_color=COLORS["border"], width=260, height=38
        )
        search.pack(side="right")
        search.pack_propagate(False)

        ctk.CTkLabel(
            search, text="", image=icon("search", COLORS["muted"], 16),
            fg_color="transparent"
        ).pack(side="left", padx=(12, 4))

        ctk.CTkEntry(
            search, textvariable=self.search_var, font=self.f_search,
            placeholder_text="Rechercher...", fg_color="transparent",
            border_width=0, text_color=COLORS["text"],
            placeholder_text_color=COLORS["muted"]
        ).pack(side="left", fill="both", expand=True, padx=(0, 8))

        self.search_var.trace_add("write", lambda *_: self.filter_rows())

        ctk.CTkFrame(table_card, fg_color=COLORS["border"], height=1).pack(fill="x")

        tree_container = ctk.CTkFrame(table_card, fg_color="transparent")
        tree_container.pack(fill="both", expand=True, padx=(1, 1), pady=(0, 8))

        tree_container.grid_rowconfigure(0, weight=1)
        tree_container.grid_columnconfigure(0, weight=1)

        self.tree = ttk.Treeview(
            tree_container, show="headings", style="Detail.Treeview"
        )
        self.tree.grid(row=0, column=0, sticky="nsew")

        vsb = ctk.CTkScrollbar(
            tree_container, orientation="vertical", command=self.tree.yview,
            width=12, fg_color="transparent",
            button_color=COLORS["track"], button_hover_color=COLORS["muted"]
        )
        vsb.grid(row=0, column=1, sticky="ns", padx=(2, 4))

        hsb = ctk.CTkScrollbar(
            tree_container, orientation="horizontal", command=self.tree.xview,
            height=12, fg_color="transparent",
            button_color=COLORS["track"], button_hover_color=COLORS["muted"]
        )
        hsb.grid(row=1, column=0, sticky="ew", padx=4, pady=(2, 2))

        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        self.tree.tag_configure("odd", background=COLORS["bg"])
        self.tree.tag_configure("even", background=COLORS["panel"])

        # ---------------------------------------------------------
        # TREND
        # ---------------------------------------------------------

        trend_card = ctk.CTkFrame(
            body, fg_color="#FFFFFF", corner_radius=16,
            border_width=1, border_color=COLORS["border"], width=390
        )
        trend_card.pack(side="right", fill="y", padx=(8, 0))
        trend_card.pack_propagate(False)

        thead2 = ctk.CTkFrame(trend_card, fg_color="transparent")
        thead2.pack(fill="x", padx=20, pady=(16, 12))

        tbox = ctk.CTkFrame(
            thead2, fg_color=COLORS["accent_soft"], width=34, height=34, corner_radius=10
        )
        tbox.pack(side="left", padx=(0, 10))
        tbox.pack_propagate(False)

        ctk.CTkLabel(
            tbox, text="", image=icon("chart", COLORS["accent"], 18),
            fg_color="transparent"
        ).pack(expand=True)

        ttl = ctk.CTkFrame(thead2, fg_color="transparent")
        ttl.pack(side="left")

        ctk.CTkLabel(
            ttl, text="Tendance", font=self.f_card_title,
            text_color=COLORS["text"], fg_color="transparent"
        ).pack(anchor="w")

        ctk.CTkLabel(
            ttl, text="7 dernières exécutions", font=self.f_small,
            text_color=COLORS["muted"], fg_color="transparent"
        ).pack(anchor="w")

        ctk.CTkFrame(trend_card, fg_color=COLORS["border"], height=1).pack(fill="x")

        # conteneur blanc : le graphique remplit tout l'espace disponible
        chart_box = ctk.CTkFrame(trend_card, fg_color="#FFFFFF", corner_radius=0)
        chart_box.pack(fill="both", expand=True, padx=8, pady=(8, 12))

        self.chart = TrendChart(chart_box, height=260)
        self.chart.pack(fill="both", expand=True)

    # =============================================================
    # SUMMARY CARD
    # =============================================================

    def create_summary(self, parent, column, value, label, icon_name,
                       icon_color, icon_bg, value_color):

        frame = ctk.CTkFrame(
            parent, fg_color=COLORS["panel"], corner_radius=16,
            border_width=1, border_color=COLORS["border"]
        )
        frame.grid(row=0, column=column, sticky="nsew",
                   padx=(0 if column == 0 else 8, 0))

        box = ctk.CTkFrame(
            frame, fg_color=icon_bg, width=44, height=44, corner_radius=12
        )
        box.pack(side="left", padx=(18, 14), pady=18)
        box.pack_propagate(False)

        ctk.CTkLabel(
            box, text="", image=icon(icon_name, icon_color, 22),
            fg_color="transparent"
        ).pack(expand=True)

        texts = ctk.CTkFrame(frame, fg_color="transparent")
        texts.pack(side="left", fill="x", expand=True, pady=14)

        ctk.CTkLabel(
            texts, text=label, font=self.f_sum_label,
            text_color=COLORS["muted"], fg_color="transparent", anchor="w"
        ).pack(anchor="w")

        ctk.CTkLabel(
            texts, text=value, font=self.f_sum_value,
            text_color=value_color, fg_color="transparent", anchor="w"
        ).pack(anchor="w")

    # =============================================================
    # DATA
    # =============================================================

    def load_data(self):

        self.rows = self.repository.get_details(
            self.card.control_id,
            self.card.date_controle,
            self.perimetre,
            self.mode_execution
        )

        self.build_dynamic_columns()

        trend = self.repository.get_trend(
            self.card.control_id,
            self.card.date_controle,
            self.perimetre,
            self.mode_execution,
            7
        )

        # get_trend renvoie les dates du plus recent au plus ancien
        # (ORDER BY date_controle DESC) -- on remet dans l'ordre
        # chronologique pour que le graphique se lise de gauche a droite
        self.chart.redraw(list(reversed(trend)))

        self.filter_rows()

    # =============================================================
    # COLONNES DYNAMIQUES JSONB
    # =============================================================

    def build_dynamic_columns(self):

        columns = [
            "cle_metier",
            "categorie"
        ]

        for row in self.rows:

            detail = row.get("detail") or {}

            for key in detail.keys():

                if key not in columns:
                    columns.append(key)

        self.columns = columns

        self.tree["columns"] = columns

        for column in columns:

            self.tree.heading(
                column,
                text=column.replace("_", " ").upper(),
                anchor="w"
            )

            if column == "cle_metier":
                width = 130

            elif column == "categorie":
                width = 170

            else:
                width = 140

            self.tree.column(
                column,
                width=width,
                minwidth=90,
                anchor="w"
            )

    # =============================================================
    # AFFICHAGE
    # =============================================================

    def filter_rows(self):

        if not hasattr(self, "rows"):
            return

        search = (
            self.search_var
            .get()
            .strip()
            .lower()
        )

        for item in self.tree.get_children():
            self.tree.delete(item)

        self.visible_values = []
        shown = 0

        for row in self.rows:

            detail = row.get("detail") or {}

            values = [
                row.get("cle_metier", ""),
                row.get("categorie", "")
            ]

            for column in self.columns[2:]:

                values.append(
                    detail.get(column, "")
                )

            searchable = " ".join(
                str(value)
                for value in values
            ).lower()

            if search and search not in searchable:
                continue

            self.tree.insert(
                "",
                "end",
                values=values,
                tags=("odd" if shown % 2 else "even",)
            )
            self.visible_values.append(values)
            shown += 1

        self.count_label.configure(
            text=f"{shown} sur {len(self.rows)} éléments — dernière exécution"
        )

    # =============================================================
    # EXPORT EXCEL
    # =============================================================

    def _default_filename(self):
        name = re.sub(r"[^\w\-]+", "_", str(self.card.control_name)).strip("_")
        date = str(self.card.date_controle or "")[:10]
        parts = [f"CTRL-{self.card.control_id:02d}", name]
        if date:
            parts.append(date)
        return "_".join(p for p in parts if p) + ".xlsx"

    @staticmethod
    def _excel_value(value):
        """Convertit les valeurs non supportées par Excel (dict, list, etc.)."""
        if value is None or isinstance(value, (int, float, bool, str)):
            return value
        if hasattr(value, "isoformat"):          # date / datetime
            return value
        if hasattr(value, "__float__"):          # Decimal
            try:
                return float(value)
            except Exception:
                pass
        return str(value)

    def export_excel(self):

        if not self.visible_values:
            messagebox.showinfo(
                "Export Excel", "Aucune ligne à exporter.", parent=self
            )
            return

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            messagebox.showerror(
                "Export Excel",
                "La bibliothèque openpyxl est requise.\n\n"
                "Installez-la avec :\n    pip install openpyxl",
                parent=self
            )
            return

        path = filedialog.asksaveasfilename(
            parent=self,
            title="Exporter vers Excel",
            defaultextension=".xlsx",
            initialfile=self._default_filename(),
            filetypes=[("Classeur Excel", "*.xlsx")]
        )

        if not path:
            return

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Détails"

            headers = [c.replace("_", " ").upper() for c in self.columns]

            head_fill = PatternFill("solid", fgColor="2F5FE0")
            head_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
            thin = Side(style="thin", color="E3E6EA")
            border = Border(left=thin, right=thin, top=thin, bottom=thin)

            ws.append(headers)
            for cell in ws[1]:
                cell.fill = head_fill
                cell.font = head_font
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.border = border
            ws.row_dimensions[1].height = 24

            for values in self.visible_values:
                ws.append([self._excel_value(v) for v in values])

            band = PatternFill("solid", fgColor="F7F8FA")
            for r_idx, row in enumerate(ws.iter_rows(min_row=2), start=0):
                for cell in row:
                    cell.border = border
                    cell.alignment = Alignment(vertical="center")
                    if r_idx % 2:
                        cell.fill = band

            # largeur automatique (bornée)
            for idx, header in enumerate(headers, start=1):
                longest = len(str(header))
                for row in self.visible_values:
                    longest = max(longest, len(str(row[idx - 1])))
                ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 3, 12), 50)

            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

            wb.save(path)

        except PermissionError:
            messagebox.showerror(
                "Export Excel",
                "Impossible d'écrire le fichier : il est peut-être ouvert dans Excel.\n"
                "Fermez-le puis réessayez.",
                parent=self
            )
            return
        except Exception as e:
            messagebox.showerror("Export Excel", f"Échec de l'export :\n{e}", parent=self)
            return

        if messagebox.askyesno(
            "Export Excel",
            f"{len(self.visible_values)} ligne(s) exportée(s).\n\nOuvrir le fichier ?",
            parent=self
        ):
            try:
                os.startfile(path)      # Windows
            except Exception:
                pass